from pathlib import Path
from typing import Dict, Any
from openpyxl import load_workbook


def load_plan_from_workbook(path: Path) -> Dict[str, Any]:
    wb = load_workbook(str(path), data_only=True)
    sheets = wb.sheetnames
    pairs: Dict[str, Dict[str, str]] = {}
    for name in sheets:
        if name.endswith('_new_version'):
            base = name[:-12]
            pairs.setdefault(base, {})['new'] = name
        else:
            pairs.setdefault(name, {})['orig'] = name

    plan: Dict[str, Any] = {}
    for base, ref in pairs.items():
        new_name = ref.get('new')
        if not new_name:
            continue
        ws = wb[new_name]
        cols = []
        data_action = ''
        notes = []
        hit_meta = False
        for r in ws.iter_rows(min_row=2, values_only=True):
            cells = list(r)
            if cells and isinstance(cells[0], str) and cells[0].strip().upper() in ("DATA:", "NOTES:"):
                hit_meta = True
            if hit_meta:
                if cells and isinstance(cells[0], str) and cells[0].strip().upper() == 'DATA:':
                    data_action = (cells[1] or '').strip().lower()
                elif cells and isinstance(cells[0], str) and cells[0].strip().upper() == 'NOTES:':
                    line = ' '.join(str(c) for c in cells[1:] if c is not None).strip()
                    if line:
                        notes.append(line)
                else:
                    line = ' '.join(str(c) for c in cells if c is not None).strip()
                    if line:
                        notes.append(line)
                continue
            if not any(c is not None and str(c).strip() != '' for c in cells):
                hit_meta = True
                continue
            if len(cells) >= 6 and cells[1]:
                col = {
                    'name': str(cells[1]).strip(),
                    'type': (str(cells[2]).strip() if cells[2] is not None else ''),
                    'notnull': bool(cells[3]) if cells[3] is not None else False,
                    'default': None if cells[4] is None else str(cells[4]),
                    'pk': bool(cells[5]) if cells[5] is not None else False,
                }
                cols.append(col)
        plan[base] = {
            'data': data_action,
            'notes': notes,
            'columns': cols,
        }
    return plan


def render_markdown(plan: Dict[str, Any]) -> str:
    lines = []
    lines.append("# Database Migration Plan (from db_schema_review.xlsx)")
    lines.append("")
    lines.append("Global rules:")
    lines.append("- IDs: switch to UUID v4 (TEXT) primary keys; generate in backend.")
    lines.append("- Foreign keys: rename to table_id (e.g., user_id, match_id).")
    lines.append("- Timestamps: rename created_at/updated_at -> created/updated; add if missing (Unix timestamp; render local time).")
    lines.append("- NOTES: if 'NO CHANGES', keep structure (still apply global rules unless explicitly contradicted).")
    lines.append("- DATA: save|delete|dump governs row retention and code removal for delete/dump.")
    lines.append("")

    for table in sorted(plan.keys()):
        p = plan[table]
        data = (p.get('data') or '').strip()
        notes = p.get('notes') or []
        cols = p.get('columns') or []
        lines.append(f"## {table}")
        lines.append(f"- DATA: {data if data else '(unspecified)'}")
        if notes:
            lines.append("- NOTES:")
            for n in notes:
                lines.append(f"  - {n}")
        else:
            lines.append("- NOTES: (none)")
        # Proposed changes per table
        lines.append("- Proposed changes:")
        # ID/PK
        has_pk = any(c.get('pk') for c in cols)
        lines.append("  - Primary key: set id TEXT (UUID v4) as PK")
        # created/updated
        colnames = {c.get('name') for c in cols}
        if 'created_at' in colnames or 'updated_at' in colnames:
            lines.append("  - Rename created_at/updated_at -> created/updated (INTEGER Unix timestamp)")
        add_created = 'created_at' not in colnames and 'created' not in colnames
        add_updated = 'updated_at' not in colnames and 'updated' not in colnames
        if add_created:
            lines.append("  - Add created (INTEGER, not null)")
        if add_updated:
            lines.append("  - Add updated (INTEGER, not null)")
        # FK rename suggestions
        fk_candidates = [c for c in cols if c.get('name') and c['name'].endswith('_id')]
        if fk_candidates:
            lines.append("  - Ensure FK names use table_id form and reference UUID TEXT")
        # Data directive
        if data in ('delete', 'dump'):
            action = 'Remove all records' if data == 'delete' else 'Drop table after migration (dump)'
            lines.append(f"  - {action}; remove relations and references in code")
        lines.append("")
    return "\n".join(lines)


def main() -> int:
    wb_path = Path('docs/db_schema_review.xlsx')
    out_path = Path('docs/db_migration_plan.md')
    if not wb_path.exists():
        print("Workbook not found: docs/db_schema_review.xlsx")
        return 1
    plan = load_plan_from_workbook(wb_path)
    md = render_markdown(plan)
    out_path.write_text(md, encoding='utf-8')
    print(f"✅ Wrote migration plan to {out_path}")
    return 0


if __name__ == '__main__':
    raise SystemExit(main())


