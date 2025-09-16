import os
import sys
import sqlite3
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet
except Exception as e:
    sys.stderr.write(f"Missing dependency openpyxl or import error: {e}\n")
    sys.exit(2)


def find_database(root: Path) -> Path | None:
    # Prefer restrike_vta.db, otherwise first .db file
    preferred = list(root.rglob("restrike_vta.db"))
    if preferred:
        return preferred[0]
    any_db = list(root.rglob("*.db"))
    return any_db[0] if any_db else None


def fetch_tables(conn: sqlite3.Connection) -> list[str]:
    cur = conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
    )
    return [r[0] for r in cur.fetchall()]


def fetch_columns(conn: sqlite3.Connection, table: str) -> list[tuple]:
    # PRAGMA table_info returns: cid, name, type, notnull, dflt_value, pk
    cur = conn.execute(f"PRAGMA table_info('{table.replace("'", "''")}')")
    return cur.fetchall()


def auto_fit_columns(ws: Worksheet):
    widths: dict[int, int] = {}
    for row in ws.iter_rows(values_only=True):
        for c, val in enumerate(row, start=1):
            text = "" if val is None else str(val)
            widths[c] = max(widths.get(c, 0), len(text))
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 10), 80)


def write_table_sheet(wb: Workbook, table: str, cols: list[tuple]):
    ws = wb.create_sheet(title=table[:31])
    headers = ["cid", "name", "type", "notnull", "dflt_value", "pk"]
    ws.append(headers)
    for row in cols:
        ws.append(list(row))
    auto_fit_columns(ws)
    ws.freeze_panes = "A2"
    return ws


def write_new_version_sheet(wb: Workbook, table: str, cols: list[tuple]):
    title = f"{table}_new_version"
    ws = wb.create_sheet(title=title[:31])
    headers = ["cid", "name", "type", "notnull", "dflt_value", "pk"]
    ws.append(headers)
    for row in cols:
        ws.append(list(row))
    # Spacer
    ws.append([])
    ws.append(["DATA:", "save | delete | DUMP"])
    ws.append(["NOTES:"])
    ws.append(["CREATED:", "unix timestamp automatically set on insert"])
    ws.append(["UPDATED:", "unix timestamp automatically set on update (if applicable)"])
    auto_fit_columns(ws)
    ws.freeze_panes = "A2"
    return ws


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    db_path = find_database(root)
    if not db_path:
        sys.stderr.write("No SQLite .db file found in repository.\n")
        return 1

    out_dir = root / "docs"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "db_schema_review.xlsx"

    conn = sqlite3.connect(str(db_path))
    try:
        tables = fetch_tables(conn)
        wb = Workbook()
        # Remove the default sheet
        default = wb.active
        wb.remove(default)

        for t in tables:
            cols = fetch_columns(conn, t)
            write_table_sheet(wb, t, cols)
            write_new_version_sheet(wb, t, cols)

        wb.save(str(out_path))
        print(f"✅ Excel generated: {out_path}")
        print(f"📦 Source DB: {db_path}")
        return 0
    finally:
        conn.close()


if __name__ == "__main__":
    raise SystemExit(main())


