import json
from pathlib import Path
from openpyxl import load_workbook


def parse_workbook(path: Path) -> dict:
    wb = load_workbook(str(path), data_only=True)
    sheets = wb.sheetnames
    pairs: dict[str, dict[str, str]] = {}
    for name in sheets:
        if name.endswith('_new_version'):
            base = name[:-12]
            pairs.setdefault(base, {})['new'] = name
        else:
            pairs.setdefault(name, {})['orig'] = name

    plan: dict[str, dict] = {}
    for base, ref in pairs.items():
        new_name = ref.get('new')
        if not new_name:
            continue
        ws = wb[new_name]
        cols = []
        data_action = ''
        notes = []
        hit_meta = False
        for r in ws.iter_rows(min_row=2, values_only=True):
            cells = list(r)
            # Meta markers
            if cells and isinstance(cells[0], str) and cells[0].strip().upper() in ("DATA:", "NOTES:"):
                hit_meta = True
            if hit_meta:
                if cells and isinstance(cells[0], str) and cells[0].strip().upper() == 'DATA:':
                    data_action = (cells[1] or '').strip().lower()
                elif cells and isinstance(cells[0], str) and cells[0].strip().upper() == 'NOTES:':
                    line = ' '.join(str(c) for c in cells[1:] if c is not None).strip()
                    if line:
                        notes.append(line)
                else:
                    line = ' '.join(str(c) for c in cells if c is not None).strip()
                    if line:
                        notes.append(line)
                continue
            if not any(c is not None and str(c).strip() != '' for c in cells):
                hit_meta = True
                continue
            if len(cells) >= 6 and cells[1]:
                col = {
                    'name': str(cells[1]).strip(),
                    'type': (str(cells[2]).strip() if cells[2] is not None else ''),
                    'notnull': bool(cells[3]) if cells[3] is not None else False,
                    'default': None if cells[4] is None else str(cells[4]),
                    'pk': bool(cells[5]) if cells[5] is not None else False,
                }
                cols.append(col)
        plan[base] = {
            'data': data_action,
            'notes': notes,
            'columns': cols,
        }
    return plan


def main() -> int:
    wb_path = Path('docs/db_schema_review.xlsx')
    if not wb_path.exists():
        print(json.dumps({'error': 'workbook_not_found', 'path': str(wb_path)}))
        return 1
    plan = parse_workbook(wb_path)
    print(json.dumps({'plan': plan}, indent=2))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())


